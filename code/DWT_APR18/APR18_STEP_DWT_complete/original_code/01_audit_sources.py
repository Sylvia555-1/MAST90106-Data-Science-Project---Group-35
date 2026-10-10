"""Audit APR18 step records before transferring synthetic DWT features.

This script only reads the source files. It does not detect change points.
Run with: /opt/anaconda3/bin/python 01_audit_sources.py
"""
from pathlib import Path
import hashlib
import json
import os

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
SOURCE = Path('/Users/jerry_t0mato/Desktop/毕业项目2/Dataset_reoh/APR18 SS SAOS STEPS LAOS UDLAOS.xlsx')
SYNTHETIC = Path('/Users/jerry_t0mato/Desktop/step_data')
OUT = ROOT / '01_source_audit'
OUT.mkdir(parents=True, exist_ok=True)
os.environ.setdefault('MPLCONFIGDIR', str(ROOT / '.plot_cache'))
os.environ.setdefault('XDG_CACHE_HOME', str(ROOT / '.plot_cache'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt


def digest(path):
    """Record hashes so the analysis can verify that sources stay unchanged."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    before = digest(SOURCE)
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    rows = list(workbook['steps'].values)
    workbook.close()
    blocks, inventory, arrays = {}, [], {}
    for index, title in enumerate(rows[0]):
        if title is None:
            continue
        column = get_column_letter(index + 1)
        nonempty = [(i + 4, r[index:index + 4]) for i, r in enumerate(rows[3:])
                    if any(value is not None for value in r[index:index + 4])]
        item = {'column': column, 'title': title, 'rows': len(nonempty)}
        if not nonempty:
            item['status'] = 'title_only_no_data'
            inventory.append(item)
            continue
        assert tuple(rows[1][index:index + 4]) == ('Step time', 'Stress', 'Strain', 'Shear rate')
        assert tuple(rows[2][index:index + 4]) == ('s', 'Pa', '%', '1/s')
        values = np.array([r[1] for r in nonempty], dtype=float)
        assert np.isfinite(values).all(), column
        assert (np.diff(values[:, 0]) > 0).all(), column
        duplicate = next((key for key, prior in arrays.items()
                          if np.array_equal(values, prior)), None)
        arrays[column] = values
        frame = pd.DataFrame(values, columns=['time_s', 'stress_Pa', 'strain_percent', 'shear_rate_1_s'])
        frame.insert(0, 'excel_row', [r[0] for r in nonempty])
        frame.to_csv(OUT / f'{column}_raw.csv', index=False)
        # Confirm exported decimals reproduce every original numeric value.
        reloaded = pd.read_csv(OUT / f'{column}_raw.csv', float_precision='round_trip')
        assert np.array_equal(reloaded.iloc[:, 1:].to_numpy(), values)
        t, stress = values[:, 0], values[:, 1]
        dt = np.diff(t)
        item.update(status='duplicate' if duplicate else 'distinct_record',
                    identical_to=duplicate,
                    source_range=f'steps!{column}4:{get_column_letter(index + 4)}68',
                    first_time_s=float(t[0]), last_time_s=float(t[-1]),
                    min_gap_s=float(dt.min()), max_gap_s=float(dt.max()),
                    points_first_0p2_s=int((t <= .2).sum()),
                    first_measured_rate=float(values[0, 3]),
                    tail_median_rate=float(np.median(values[-10:, 3])),
                    observed_min_time_s=float(t[stress.argmin()]),
                    observed_max_time_s=float(t[stress.argmax()]))
        inventory.append(item)
        blocks[column] = frame
    pd.DataFrame(inventory).to_csv(OUT / 'step_inventory.csv', index=False)

    # Inspect observation support without constructing additional measurements.
    support = []
    for column, frame in blocks.items():
        t = frame.time_s.to_numpy()
        for left in range(1, 63):
            right = left + 1
            crossing = (t[:-1] < right) & (t[1:] > left)
            support.append(dict(column=column, left_s=left, right_s=right,
                                original_points=int(((t >= left) & (t < right)).sum()),
                                maximum_crossing_gap_s=float(np.diff(t)[crossing].max())))
    support = pd.DataFrame(support)
    support.to_csv(OUT / 'original_support_per_second.csv', index=False)

    # A 100 Hz grid is a diagnostic, not an accepted replacement for raw data.
    fidelity = []
    selected = blocks['BH']
    t = selected.time_s.to_numpy()
    for dt in (.01, .0003125, .00015625):
        grid = t[0] + np.arange(int(np.floor((t[-1] - t[0]) / dt)) + 1) * dt
        for channel in ('stress_Pa', 'shear_rate_1_s'):
            original = selected[channel].to_numpy()
            interpolated = np.interp(grid, t, original)
            # Do not extrapolate the diagnostic round-trip beyond the last grid node.
            valid = t <= grid[-1]
            error = np.max(np.abs(np.interp(t[valid], grid, interpolated) - original[valid]))
            fidelity.append(dict(dt_s=dt, channel=channel,
                                 max_roundtrip_error=float(error),
                                 fraction_of_original_range=float(error / np.ptp(original)),
                                 raw_knots_in_roundtrip=int(valid.sum()),
                                 all_raw_points_used_as_interpolation_knots=True))
    pd.DataFrame(fidelity).to_csv(OUT / 'grid_fidelity_diagnostic.csv', index=False)

    synthetic_profile = []
    hashes = {str(SOURCE): before, str(SYNTHETIC / 'manifest.json'): digest(SYNTHETIC / 'manifest.json')}
    manifest = json.loads((SYNTHETIC / 'manifest.json').read_text())
    for relative in ('test/control_1.csv', 'test/test_clot_1.csv', 'train/cal_02_0p1_to_1.csv'):
        path = SYNTHETIC / relative
        hashes[str(path)] = digest(path)
        frame = pd.read_csv(path)
        st = frame['time_s'].to_numpy()
        declared = next(item for item in manifest['records'] if item['file'] == path.name)
        synthetic_profile.append(dict(file=relative, rows=len(frame), columns=list(frame.columns),
                                      start_s=float(st[0]), end_s=float(st[-1]),
                                      uniform_100_Hz=bool(np.allclose(np.diff(st), .01)),
                                      declared_pre_rate_1_s=declared['pre_rate_1_s'],
                                      declared_target_rate_1_s=declared['target_rate_1_s'],
                                      declared_added_onset_s=declared['clot_onset_s']))
    assert all(digest(Path(path)) == h for path, h in hashes.items())
    assert np.array_equal(arrays['F'], arrays['AS']) and np.array_equal(arrays['AX'], arrays['AS'])
    summary = dict(source_hashes=hashes, inventory=inventory,
                   distinct_records=sum(x['status'] == 'distinct_record' for x in inventory),
                   duplicate_group=['F', 'AS', 'AX'],
                   duplicate_warning='F and AX titles conflict with the observed 5 to 0.5/s input.',
                   synthetic_profile=synthetic_profile,
                   BH_complete_one_second_windows=len(support[support.column == 'BH']),
                   BH_windows_with_no_original_point=int(((support.column == 'BH') & (support.original_points == 0)).sum()),
                   grid_diagnostics=fidelity,
                   checks={'source_hashes_unchanged': True, 'exports_exact': True,
                           'finite_values': True, 'strictly_increasing_time': True},
                   scope='Source audit only. No DWT features or change-point detection yet.')
    (OUT / 'audit.json').write_text(json.dumps(summary, indent=2))

    fig, axes = plt.subplots(3, 1, figsize=(11, 8), constrained_layout=True)
    for axis, channel, unit in zip(axes[:2], ('stress_Pa', 'shear_rate_1_s'), ('Stress (Pa)', 'Shear rate (1/s)')):
        for column in ('BC', 'BH', 'BM', 'BR', 'AI', 'AN', 'AS'):
            frame = blocks[column]
            item = next(r for r in inventory if r['column'] == column)
            axis.plot(frame.time_s, frame[channel], '.-', ms=3, lw=.8,
                      label=f"{column}: {item['first_measured_rate']:.1g} to {item['tail_median_rate']:.2g}/s")
        axis.set_xscale('log')
        axis.set_ylabel(unit)
        axis.grid(alpha=.2)
    axes[0].legend(ncol=4, fontsize=8)
    selected_support = support[support.column == 'BH']
    axes[2].bar(selected_support.left_s + .5, selected_support.original_points, width=.8)
    axes[2].set(xlabel='Recorded time (s)', ylabel='Original points per 1 s',
                title='BH observation support: an empty window is not a measured constant signal')
    fig.suptitle('APR18 step inventory and sampling support\nOriginal samples; connecting lines are visual guides')
    fig.savefig(OUT / 'raw_inventory_and_support.png', dpi=160)
    plt.close(fig)
    print(json.dumps({k:summary[k] for k in ('distinct_records','BH_complete_one_second_windows',
          'BH_windows_with_no_original_point','synthetic_profile','grid_diagnostics','checks')}, indent=2))


if __name__ == '__main__':
    main()
